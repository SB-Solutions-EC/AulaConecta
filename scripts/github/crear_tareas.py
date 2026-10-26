"""Crea en GitHub una tarea (issue) por cada subactividad del cronograma.

Lee docs/planificacion/Planificacion.xlsx (una hoja por integrante) y, por cada día de trabajo:
  - crea el issue "T<MMDD>-<inicial> · <subactividad>" asignado al integrante,
  - lo pone en el milestone de su sprint y con la etiqueta "tarea",
  - incluye actividad del día, entregable, evidencia y carpeta del Drive,
  - cierra como completadas las tareas anteriores a la fecha de corte.
Se puede ejecutar varias veces: las tareas que ya existen no se duplican.

Uso:  uv run --with openpyxl python scripts/github/crear_tareas.py [--corte AAAA-MM-DD]
Requiere gh autenticado con permiso "repo".
"""
import argparse
import json
import subprocess
import sys
import time
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO = "SB-Solutions-EC/AulaConecta"
XLSX = Path(__file__).resolve().parents[2] / "docs" / "planificacion" / "Planificacion.xlsx"
PERSONAS = {"Marcelo": ("M", "cmacuna7"), "Abner": ("A", "adarboleda"), "Christian": ("C", "cmbonifaz")}
SPRINTS = [
    (1, "2026-09-30", "2026-10-06", "Sprint 1 – Concepción y propuesta del SI", "Tema del proyecto y Entregable 1"),
    (2, "2026-10-07", "2026-10-20", "Sprint 2 – Procesos, requisitos y análisis", "SRS, BPMN, casos de uso, riesgos y Entregable 2"),
    (3, "2026-10-21", "2026-10-30", "Sprint 3 – Repositorio y CI/CD básico", ""),
    (4, "2026-11-04", "2026-11-13", "Sprint 4 – Autenticación y matrícula", ""),
    (5, "2026-11-16", "2026-11-27", "Sprint 5 – Calificaciones", ""),
    (6, "2026-11-30", "2026-12-08", "Sprint 6 – Reportes y tableros", ""),
    (7, "2026-12-09", "2026-12-18", "Sprint 7 – Planificación curricular", ""),
    (8, "2026-12-21", "2026-12-29", "Sprint 8 – Notificaciones por WhatsApp", ""),
    (9, "2026-12-30", "2027-01-08", "Sprint 9 – Trabajo sin conexión", ""),
    (10, "2027-01-11", "2027-01-15", "Sprint 10 – Aceptación y entrega", ""),
]


def gh(*args, data=None):
    for intento in range(5):
        r = subprocess.run(["gh", *args], input=data, capture_output=True, text=True, encoding="utf-8")
        if r.returncode == 0:
            return r.stdout
        if "rate limit" in r.stderr.lower() or "secondary" in r.stderr.lower():
            espera = 60 * (intento + 1)
            print(f"  límite de GitHub, esperando {espera} s…", flush=True)
            time.sleep(espera)
            continue
        raise RuntimeError(f"gh {' '.join(args[:3])}…: {r.stderr.strip()}")
    raise RuntimeError("Se agotaron los reintentos por límite de GitHub")


def sprint_de(iso):
    for s in SPRINTS:
        if s[1] <= iso <= s[2]:
            return s
    return None


def asegurar_milestones():
    existentes = {m["title"]: m["number"] for m in json.loads(gh("api", f"repos/{REPO}/milestones?state=all&per_page=100"))}
    for n, ini, fin, titulo, desc in SPRINTS:
        if titulo not in existentes:
            payload = json.dumps({"title": titulo, "due_on": f"{fin}T23:59:59Z", "description": desc})
            m = json.loads(gh("api", f"repos/{REPO}/milestones", "--input", "-", data=payload))
            existentes[titulo] = m["number"]
            print(f"  milestone creado: {titulo}")
    return existentes


def tareas():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for hoja, (ini, usuario) in PERSONAS.items():
        ws = wb[hoja]
        for r in range(3, ws.max_row + 1):
            f = ws.cell(r, 2).value
            if not isinstance(f, datetime) or not ws.cell(r, 4).value:
                continue
            d = f.date()
            yield {
                "fecha": d, "hoja": hoja, "usuario": usuario,
                "clave": f"T{d:%m%d}-{ini}",
                "actividad": ws.cell(r, 5).value, "sub": ws.cell(r, 6).value,
                "entregable": ws.cell(r, 7).value, "evidencia": ws.cell(r, 8).value,
                "carpeta": ws.cell(r, 9).value or "",
            }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--corte", default="2026-10-21", help="Las tareas anteriores a esta fecha se cierran como completadas")
    args = ap.parse_args()
    corte = date.fromisoformat(args.corte)

    gh("label", "create", "tarea", "--repo", REPO, "--color", "BFD4F2", "--description", "Subactividad diaria del cronograma", "--force")
    ms = asegurar_milestones()
    existentes = {}
    for it in json.loads(gh("issue", "list", "--repo", REPO, "--state", "all", "--limit", "1000", "--label", "tarea", "--json", "number,title,state")):
        existentes[it["title"].split(" · ")[0]] = it

    creadas = cerradas = 0
    for t in tareas():
        s = sprint_de(t["fecha"].isoformat())
        sub = t["sub"].strip().rstrip(".")
        titulo = f"{t['clave']} · {sub[:110] + '…' if len(sub) > 110 else sub}"
        cuerpo = (
            f"**Fecha:** {t['fecha']:%d/%m/%Y} · **Integrante:** {t['hoja']} (@{t['usuario']}) · **Sprint {s[0]}**\n\n"
            f"**Actividad del día:** {t['actividad']}\n\n"
            f"**Subactividad:** {t['sub']}\n\n"
            f"**Entregable:** {t['entregable']}\n\n"
            f"**Evidencia:** {t['evidencia']}\n\n"
            f"**Carpeta en Drive:** `{t['carpeta']}`\n\n"
            "_Generada desde docs/planificacion/Planificacion.xlsx._"
        )
        previa = existentes.get(t["clave"])
        if previa:
            numero, estado = previa["number"], previa["state"]
        else:
            url = gh("issue", "create", "--repo", REPO, "--title", titulo, "--body", cuerpo,
                     "--label", "tarea", "--assignee", t["usuario"], "--milestone", s[3]).strip()
            numero, estado = int(url.rsplit("/", 1)[1]), "OPEN"
            creadas += 1
            print(f"  #{numero} {t['clave']}", flush=True)
            time.sleep(1)  # evita el límite secundario de GitHub
        if t["fecha"] < corte and estado == "OPEN":
            gh("issue", "close", str(numero), "--repo", REPO, "--reason", "completed")
            cerradas += 1
            time.sleep(0.5)

    # Los milestones de sprints ya terminados se cierran
    for n, ini, fin, titulo, _ in SPRINTS:
        if date.fromisoformat(fin) < corte:
            gh("api", "-X", "PATCH", f"repos/{REPO}/milestones/{ms[titulo]}", "-f", "state=closed")
    print(f"Listo: {creadas} tareas creadas, {cerradas} cerradas como completadas.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
