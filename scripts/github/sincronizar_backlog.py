"""Sincroniza el Product Backlog de AulaConecta en GitHub con docs/planificacion/Planificacion_v2.xlsx.

Qué hace:
  - etiquetas y un milestone por sprint (1 a 10) con el objetivo de la hoja "Sprints"
  - ítems del Product Backlog: 15 historias de usuario (HU), habilitadores técnicos (TEC) y entregables (DOC)
  - una tarea por integrante y día (T{MMDD}-{M|A|C}) como sub-issue de su ítem del backlog
  - el tablero "AulaConecta – Product Backlog" en GitHub Projects con los campos Status, Sprint,
    Tipo, Área, Prioridad, Wiegers, Requisitos, Fecha y Horas

Se puede ejecutar más de una vez: actualiza lo que existe y no duplica. El Status de un ítem solo se
asigna la primera vez que entra al tablero, para no pisar lo que el equipo mueve a mano.

Uso:  python scripts/github/sincronizar_backlog.py [--simular]
Requisitos: openpyxl y gh autenticado con los permisos repo y project (gh auth refresh -s project).
"""
import json
import subprocess
import sys
import time
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ORG = "SB-Solutions-EC"
NOMBRE_REPO = "AulaConecta"
REPO = f"{ORG}/{NOMBRE_REPO}"
TITULO_PROYECTO = "AulaConecta – Product Backlog"
EXCEL = Path(__file__).resolve().parents[2] / "docs" / "planificacion" / "Planificacion_v2.xlsx"
SIMULAR = "--simular" in sys.argv

INTEGRANTES = {  # hoja del Excel -> (letra del código, usuario de GitHub)
    "Marcelo": ("M", "cmacuna7"),
    "Abner": ("A", "adarboleda"),
    "Christian": ("C", "cmbonifaz"),
}
EQUIPO = ["cmacuna7", "adarboleda", "cmbonifaz"]

# Títulos cortos de los milestones; el objetivo completo va en la descripción
NOMBRE_SPRINT = {
    1: "Concepción y propuesta del SI",
    2: "Procesos, requisitos y análisis",
    3: "Arquitectura de microservicios y CI/CD básico",
    4: "Imágenes en el CI, autenticación y matrícula",
    5: "Ambiente de pruebas y calificaciones",
    6: "Reportes, tableros y monitoreo",
    7: "Planificación curricular",
    8: "Notificaciones por WhatsApp",
    9: "Trabajo sin conexión y rendimiento",
    10: "Aceptación y entrega",
}

ETIQUETAS = [
    ("historia", "1D76DB", "Historia de usuario del Product Backlog"),
    ("habilitador", "6F42C1", "Trabajo técnico del Product Backlog que no es una historia de usuario"),
    ("entregable", "0E8A16", "Entregable académico de la asignatura"),
    ("tarea", "BFD4F2", "Tarea diaria de un integrante (Sprint Backlog)"),
    ("prioridad:alta", "B60205", "Necesaria en la primera versión"),
    ("prioridad:media", "FBCA04", "Deseable en la primera versión"),
    ("opcional", "C5DEF5", "Se hace solo si el tiempo alcanza"),
    ("area:ms-identidad", "D4C5F9", "Microservicio de identidad (Java/Spring Boot)"),
    ("area:ms-academico", "5319E7", "Microservicio académico (NestJS/TypeScript) y cálculo de notas"),
    ("area:ms-notificaciones", "F9D0C4", "Microservicio de notificaciones (Python/FastAPI)"),
    ("area:pwa", "0E8A16", "PWA React y trabajo sin conexión"),
    ("area:devsecops", "D93F0B", "CI/CD, seguridad, despliegue y monitoreo"),
    ("area:docs", "C2E0C6", "Documentación y entregables"),
    ("defecto", "E11D21", "Error encontrado en pruebas o en la revisión"),
]
ETIQUETAS_OBSOLETAS = ["area:backend", "area:frontend"]  # de la v1 (monolito)

AREA_CAMPO = {"ms-identidad": "ms-identidad", "ms-academico": "ms-académico",
              "ms-notificaciones": "ms-notificaciones", "pwa": "PWA", "devsecops": "DevSecOps", "docs": "Documentación"}

# Ítems del Product Backlog:
# código, tipo, título, requisitos, prioridad, puesto en Wiegers, sprint (None = opcional), área, responsables, descripción
BACKLOG = [
    ("DOC-01", "entregable", "Entregable 1: Propuesta del sistema de información", "—", "alta", None, 1, "docs", EQUIPO,
     "Tema y cliente, problemática, matriz SDLC y relación necesidad → requisito → artefacto."),
    ("DOC-02", "entregable", "Entregable 2: Procesos, requisitos e involucrados", "RF-01 a RF-38, RNF", "alta", None, 2, "docs", EQUIPO,
     "Proceso actual y propuesto, BPMN en Bizagi, SRS v0.1, validación del prototipo, taller de riesgos, alcance, stakeholders y casos de uso."),
    ("TEC-01", "habilitador", "Organización, repositorio, ramas protegidas y Product Backlog", "RNF-12 a RNF-15", "alta", None, 3, "devsecops", EQUIPO,
     "Organización con los tres integrantes como owners, repositorio con main/test/dev, CODEOWNERS, reglas de protección, secretos y Product Backlog en GitHub Projects."),
    ("DOC-03", "entregable", "Entregable 3: Metodología, planificación, tecnologías y prototipo", "—", "alta", None, 3, "docs", EQUIPO,
     "Selección de tecnologías, prototipo anonimizado revisado con Gitleaks y entrega del Entregable 3."),
    ("TEC-02", "habilitador", "Arquitectura de 3 microservicios, modelo de datos y modelo de amenazas", "RNF-01, RNF-04", "alta", None, 3, "devsecops", EQUIPO,
     "Modelo de datos por servicio, STRIDE, C4 (contexto y contenedores), diseño offline y contratos OpenAPI v0 del monorepo."),
    ("TEC-03", "habilitador", "Esqueletos de los microservicios, gateway, PWA, bases de datos y RabbitMQ", "—", "alta", None, 3, "devsecops", EQUIPO,
     "ms-academico, ms-identidad, ms-notificaciones, gateway Traefik y PWA con /health; una base por servicio con migraciones; RabbitMQ y docker-compose."),
    ("HU-15", "historia", "Como equipo queremos un pipeline con pruebas y análisis de seguridad que avise por Telegram", "RNF-12 a RNF-15", "alta", None, 3, "devsecops", ["adarboleda"], ""),
    ("HU-01", "historia", "Como usuario quiero iniciar sesión y ver solo las opciones de mi rol", "RF-01 a RF-03, RNF-12", "alta", 1, 4, "ms-identidad", ["adarboleda"], ""),
    ("HU-03", "historia", "Como administrador quiero gestionar docentes, cursos, materias y año lectivo", "RF-05 a RF-07 (RF-08 opcional)", "alta", 2, 4, "ms-academico", ["cmacuna7"], ""),
    ("HU-02", "historia", "Como secretaría quiero matricular estudiantes e importar la nómina desde Excel", "RF-04, RF-29", "alta", 6, 4, "ms-academico", ["cmacuna7", "adarboleda"], ""),
    ("TEC-04", "habilitador", "Ambiente de pruebas con despliegue continuo, datos de prueba y pruebas e2e", "RNF-13", "alta", None, 5, "devsecops", ["adarboleda"],
     "CD de la rama test con GitHub Environments, seed de datos ficticios y prueba e2e de login con Playwright."),
    ("HU-07", "historia", "Como docente quiero registrar los aportes y que la nota final se calcule igual que la hoja oficial", "RF-17, RF-18, RF-27, RNF-09, RNF-10", "alta", 5, 5, "ms-academico", ["cmacuna7"], ""),
    ("HU-08", "historia", "Como docente quiero crear tareas y registrar asistencia y comportamiento", "RF-19 a RF-22", "alta", 8, 5, "ms-academico", ["cmbonifaz"], ""),
    ("HU-11", "historia", "Como estudiante o representante quiero ver notas, asistencia, horario y alertas", "RF-28, RF-37, RF-38", "alta", 3, 6, "ms-academico", ["cmbonifaz"], ""),
    ("HU-09", "historia", "Como docente quiero obtener la hoja de estándares y la valoración de logros sin cálculos manuales", "RF-23, RF-24, RF-26", "alta", 4, 6, "ms-academico", ["cmacuna7"], ""),
    ("HU-10", "historia", "Como vicerrector quiero ver un tablero con el rendimiento de todos los cursos", "RF-25", "alta", 4, 6, "ms-academico", ["cmbonifaz"], ""),
    ("TEC-05", "habilitador", "Caché y monitoreo con Redis, Prometheus y Grafana", "RNF-06", "alta", None, 6, "devsecops", ["adarboleda"],
     "Redis como caché del tablero y métricas de los 3 servicios en Prometheus + Grafana."),
    ("HU-05", "historia", "Como vicerrector quiero aprobar o devolver la planificación con observaciones por campo", "RF-10, RF-12 a RF-14", "alta", 7, 7, "ms-academico", ["cmacuna7"], ""),
    ("HU-06", "historia", "Como docente quiero exportar la planificación aprobada a Word y PDF y ver su calendario", "RF-15, RF-16", "alta", 9, 7, "ms-academico", ["cmbonifaz"], ""),
    ("HU-04", "historia", "Como docente quiero llenar la planificación en la plantilla del Ministerio y enviarla a revisión", "RF-09, RF-11", "alta", 10, 7, "ms-academico", ["cmbonifaz"], ""),
    ("HU-14", "historia", "Como representante quiero recibir un WhatsApp cuando mi representado baje de 7 o falte", "RF-29 a RF-35, RNF-05", "alta", 13, 8, "ms-notificaciones", ["cmacuna7", "adarboleda"], ""),
    ("HU-13", "historia", "Como docente quiero registrar notas y asistencia sin internet y que se sincronicen después", "RNF-01, RF-35", "alta", 15, 9, "pwa", ["cmbonifaz", "adarboleda"], ""),
    ("TEC-06", "habilitador", "Pruebas de carga con 400 estudiantes y 60 docentes", "RNF-06", "alta", None, 9, "devsecops", ["adarboleda"],
     "Línea base con k6 (sprint 5), carga del tablero del estudiante (sprint 6) y escenarios de carga, pico, estrés y resistencia con optimización (sprint 9)."),
    ("TEC-07", "habilitador", "Pruebas de aceptación con los criterios del SRS y corrección de hallazgos", "Criterios de aceptación del SRS", "alta", None, 10, "devsecops", EQUIPO,
     "Aceptación de cálculos y flujos, informe de seguridad final y corrección de hallazgos."),
    ("TEC-08", "habilitador", "Despliegue a producción, respaldos y release v1.0.0", "RNF-08", "alta", None, 10, "devsecops", ["adarboleda"],
     "Respaldo de las 3 bases con prueba de restauración y release v1.0.0 con aprobación manual en main."),
    ("TEC-09", "habilitador", "Documentación, capacitación y presentación final", "RNF-11", "alta", None, 10, "docs", EQUIPO,
     "Manual de usuario, documentación técnica, manual de operación, material de capacitación y presentación final."),
    ("HU-12", "historia", "Como docente quiero mensajería y recursos de estudio por materia", "RF-36", "media", 14, None, "pwa", ["cmbonifaz"], ""),
]

# A qué ítem del backlog pertenece cada tarea, por fecha (MMDD) y con excepciones por código.
# Las fechas sin ítem son los cierres de sprint (integración, demostración e informe).
ITEM_POR_FECHA = {
    **dict.fromkeys(["0930", "1001", "1002", "1005", "1006"], "DOC-01"),
    **dict.fromkeys(["1007", "1008", "1012", "1013", "1014", "1015", "1016", "1019", "1020"], "DOC-02"),
    "1021": "TEC-01", "1022": "DOC-03", "1023": "TEC-02", "1026": "TEC-02",
    "1027": "TEC-03", "1028": "TEC-03", "1029": "HU-15", "1104": "HU-15",
    "1105": "HU-01", "1106": "HU-01", "1109": "HU-03", "1110": "HU-03", "1111": "HU-02", "1112": "HU-02",
    "1116": "TEC-04", "1117": "HU-07", "1118": "HU-07", "1119": "HU-08", "1120": "HU-08",
    "1123": "HU-08", "1124": "HU-08", "1125": "HU-07", "1126": "HU-07",
    "1130": "HU-09", "1201": "HU-09", "1202": "HU-10", "1203": "HU-09", "1204": "HU-11", "1207": "HU-11",
    "1209": "HU-04", "1210": "HU-04", "1211": "HU-04", "1214": "HU-05", "1215": "HU-05", "1216": "HU-06", "1217": "HU-06",
    **dict.fromkeys(["1221", "1222", "1223", "1224", "1228"], "HU-14"),
    **dict.fromkeys(["1230", "1231", "0104", "0105", "0106"], "HU-13"),
    "0107": "TEC-06", "0111": "TEC-07", "0112": "TEC-07", "0113": "TEC-08", "0114": "TEC-09", "0115": "TEC-09",
}
ITEM_POR_CODIGO = {
    "T1030-M": "TEC-03", "T1030-A": "HU-15", "T1104-C": "HU-01", "T1126-A": "TEC-06",
    "T1202-A": "TEC-05", "T1204-A": "TEC-06", "T0113-C": "TEC-09",
}

DEFINICION_TERMINADO = ("código en `dev`, pruebas y CI en verde, sin vulnerabilidades críticas o altas, "
                        "evidencia registrada y aceptada en la revisión del sprint.")

ESTADOS = [  # opciones del campo Status
    ("Backlog", "GRAY", "Product Backlog: aún no entra a un sprint"),
    ("Por hacer", "BLUE", "Sprint Backlog: comprometido en el sprint actual"),
    ("En progreso", "YELLOW", "Alguien está trabajando en esto"),
    ("En revisión", "PURPLE", "PR abierto o esperando la revisión del sprint"),
    ("Terminado", "GREEN", "Cumple la definición de terminado"),
]
TIPOS = [("Historia de usuario", "BLUE", "HU del Product Backlog"),
         ("Habilitador técnico", "PURPLE", "Trabajo técnico del Product Backlog"),
         ("Entregable", "GREEN", "Entregable académico"),
         ("Tarea", "GRAY", "Tarea diaria de un integrante")]
TIPO_CAMPO = {"historia": "Historia de usuario", "habilitador": "Habilitador técnico", "entregable": "Entregable"}
PRIORIDADES = [("Alta", "RED", "Necesaria en la primera versión"),
               ("Media", "YELLOW", "Deseable en la primera versión"),
               ("Opcional", "GRAY", "Solo si el tiempo alcanza")]
AREAS = [("ms-identidad", "PINK", "Java/Spring Boot"), ("ms-académico", "PURPLE", "NestJS/TypeScript"),
         ("ms-notificaciones", "ORANGE", "Python/FastAPI"), ("PWA", "GREEN", "React + Vite"),
         ("DevSecOps", "RED", "CI/CD, seguridad y despliegue"), ("Documentación", "GRAY", "Entregables y manuales")]


# ---------------------------------------------------------------- gh

def gh(args, entrada=None):
    for intento in range(6):
        r = subprocess.run(["gh", *args], input=json.dumps(entrada) if entrada is not None else None,
                           capture_output=True, text=True, encoding="utf-8")
        if r.returncode == 0:
            return json.loads(r.stdout) if r.stdout.strip() else None
        error = r.stderr + r.stdout
        if any(s in error.lower() for s in ("rate limit", "abuse", "502", "503", "timeout", "was submitted too quickly",
                                                 "something went wrong")):
            espera = 30 * (intento + 1)
            print(f"    (límite de GitHub, espero {espera}s)")
            time.sleep(espera)
            continue
        raise RuntimeError(f"gh {' '.join(args[:3])}: {error.strip()}")
    raise RuntimeError("GitHub siguió rechazando la petición")


def gql(consulta, **variables):
    r = gh(["api", "graphql", "--input", "-"], {"query": consulta, "variables": variables})
    if r.get("errors"):
        raise RuntimeError(json.dumps(r["errors"], ensure_ascii=False))
    return r["data"]


def rest(metodo, ruta, datos=None):
    return gh(["api", "--method", metodo, ruta, "--input", "-"], datos or {})


def mutaciones_en_lote(mutaciones, tam=20):
    """Envía varias mutaciones en una sola petición (alias m0, m1, ...)."""
    for i in range(0, len(mutaciones), tam):
        lote = mutaciones[i:i + tam]
        gql("mutation {\n" + "\n".join(f"m{j}: {m}" for j, m in enumerate(lote)) + "\n}")
        time.sleep(1)


def q(texto):
    return json.dumps(texto, ensure_ascii=False)


# ---------------------------------------------------------------- Excel

def leer_excel():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    sprints = {}
    for fila in wb["Sprints"].iter_rows(min_row=2, values_only=True):
        nombre, inicio, fin, objetivo, hito = fila[1:6]
        if isinstance(nombre, str) and nombre.startswith("Sprint "):
            sprints[int(nombre.split()[1])] = {"inicio": inicio.date(), "fin": fin.date(), "objetivo": objetivo, "hito": hito}
    tareas = []
    for hoja, (letra, usuario) in INTEGRANTES.items():
        for f in wb[hoja].iter_rows(min_row=3, values_only=True):
            fecha, _, horas, actividad, sub, entregable, evidencia, carpeta = f[1:9]
            if not hasattr(fecha, "date") or not sub:
                continue
            fecha = fecha.date()
            sprint = next(n for n, s in sprints.items() if s["inicio"] <= fecha <= s["fin"])
            codigo = f"T{fecha:%m%d}-{letra}"
            tareas.append({
                "codigo": codigo, "fecha": fecha, "horas": horas, "sprint": sprint, "usuario": usuario,
                "item": ITEM_POR_CODIGO.get(codigo, ITEM_POR_FECHA.get(f"{fecha:%m%d}")),
                "titulo": f"{codigo} · {sub}",
                "cuerpo": (f"**Fecha:** {fecha:%d/%m/%Y} · **Integrante:** {hoja} (@{usuario}) · **Sprint {sprint}**\n\n"
                           f"**Actividad del día:** {actividad}\n\n**Subactividad:** {sub}\n\n"
                           f"**Entregable:** {entregable}\n\n**Evidencia:** {evidencia}\n\n"
                           f"**Carpeta en Drive:** `{carpeta}`\n\n**Horas:** {horas}\n\n"
                           f"_Generada desde docs/planificacion/Planificacion_v2.xlsx._"),
            })
    return sprints, tareas


def cuerpo_item(item, horas):
    codigo, tipo, titulo, reqs, prio, wiegers, sprint, area, _, desc = item
    lineas = []
    if tipo == "historia":
        lineas.append(f"**Historia:** {titulo}")
    elif desc:
        lineas.append(desc)
    lineas.append(f"**Requisitos:** {reqs}")
    puesto = f" · puesto {wiegers} en la matriz de Wiegers" if wiegers else ""
    lineas.append(f"**Prioridad:** {prio}{puesto}")
    lineas.append(f"**Sprint planificado:** {f'Sprint {sprint} – {NOMBRE_SPRINT[sprint]}' if sprint else 'opcional (si el tiempo alcanza)'}")
    lineas.append(f"**Área:** {AREA_CAMPO[area]} · **Estimación:** {horas} h (suma de sus tareas)")
    lineas.append("**Tareas:** ver los sub-issues de este ítem (una por integrante y día).")
    if tipo != "entregable":
        lineas.append("**Criterios de aceptación:** los de los requisitos indicados en el SRS (`docs/requisitos`).")
        lineas.append(f"**Definición de terminado:** {DEFINICION_TERMINADO}")
    lineas.append("_Generado desde docs/planificacion/Planificacion_v2.xlsx._")
    return "\n\n".join(lineas)


# ---------------------------------------------------------------- GitHub: repositorio

def issues_existentes():
    issues, cursor = {}, None
    while True:
        d = gql("""query($o:String!,$n:String!,$c:String){repository(owner:$o,name:$n){
            issues(first:100,after:$c){pageInfo{hasNextPage endCursor} nodes{
              id number title body state milestone{number} parent{id}
              assignees(first:10){nodes{login}} labels(first:20){nodes{name}}}}}}""",
                o=ORG, n=NOMBRE_REPO, c=cursor)["repository"]["issues"]
        for i in d["nodes"]:
            issues[i["title"].split(" · ")[0]] = i
        if not d["pageInfo"]["hasNextPage"]:
            return issues
        cursor = d["pageInfo"]["endCursor"]


def sincronizar_etiquetas():
    print("== Etiquetas")
    for nombre, color, desc in ETIQUETAS:
        if not SIMULAR:
            gh(["label", "create", nombre, "--repo", REPO, "--color", color, "--description", desc, "--force"])
    existentes = {e["name"] for e in gh(["label", "list", "--repo", REPO, "--limit", "200", "--json", "name"])}
    for nombre in ETIQUETAS_OBSOLETAS:
        if nombre in existentes:
            print(f"  elimino etiqueta obsoleta {nombre}")
            if not SIMULAR:
                gh(["label", "delete", nombre, "--repo", REPO, "--yes"])


def sincronizar_milestones(sprints):
    print("== Milestones")
    existentes = {int(m["title"].split()[1]): m for m in rest("GET", f"repos/{REPO}/milestones?state=all&per_page=100")}
    numeros = {}
    for n, s in sprints.items():
        datos = {"title": f"Sprint {n} – {NOMBRE_SPRINT[n]}",
                 "description": f"{s['objetivo']}. Hito: {s['hito']} ({s['inicio']:%d/%m/%Y} – {s['fin']:%d/%m/%Y})",
                 "due_on": f"{s['fin']}T23:59:59Z", "state": "closed" if s["fin"] < date.today() else "open"}
        m = existentes.get(n)
        if m and all(m.get(k) == v for k, v in datos.items() if k != "due_on") and m["due_on"].startswith(str(s["fin"])):
            numeros[n] = m["number"]
            continue
        print(f"  {'actualizo' if m else 'creo'} {datos['title']}")
        if SIMULAR:
            numeros[n] = m["number"] if m else 0
        elif m:
            numeros[n] = rest("PATCH", f"repos/{REPO}/milestones/{m['number']}", datos)["number"]
        else:
            numeros[n] = rest("POST", f"repos/{REPO}/milestones", datos)["number"]
    return numeros


def asegurar_issue(existentes, codigo, titulo, cuerpo, etiquetas, asignados, milestone, cerrado=None):
    """Crea o actualiza un issue y devuelve su id de nodo."""
    datos = {"title": titulo, "body": cuerpo, "labels": sorted(etiquetas), "assignees": sorted(asignados), "milestone": milestone}
    i = existentes.get(codigo)
    if i is None:
        print(f"  creo {titulo[:90]}")
        if SIMULAR:
            return None
        nuevo = rest("POST", f"repos/{REPO}/issues", datos)
        if cerrado:
            rest("PATCH", f"repos/{REPO}/issues/{nuevo['number']}", {"state": "closed", "state_reason": "completed"})
        time.sleep(1)
        return nuevo["node_id"]
    actual = {"title": i["title"], "body": i["body"], "labels": sorted(e["name"] for e in i["labels"]["nodes"]),
              "assignees": sorted(a["login"] for a in i["assignees"]["nodes"]),
              "milestone": (i["milestone"] or {}).get("number")}
    cambios = [k for k in datos if datos[k] != actual[k]]
    if cambios:
        print(f"  actualizo #{i['number']} {codigo} ({', '.join(cambios)})")
        if not SIMULAR:
            rest("PATCH", f"repos/{REPO}/issues/{i['number']}", {k: datos[k] for k in cambios})
            time.sleep(0.5)
    return i["id"]


# ---------------------------------------------------------------- GitHub Projects

CAMPOS_PROYECTO = """
  ... on ProjectV2FieldCommon { id name dataType }
  ... on ProjectV2SingleSelectField { options { id name } }
  ... on ProjectV2IterationField { configuration { iterations { id title startDate } completedIterations { id title startDate } } }
"""


def obtener_proyecto():
    d = gql("""query($o:String!,$n:String!){organization(login:$o){id projectsV2(first:50){nodes{id number title url}}}
               repository(owner:$o,name:$n){id}}""", o=ORG, n=NOMBRE_REPO)
    for p in d["organization"]["projectsV2"]["nodes"]:
        if p["title"] == TITULO_PROYECTO:
            return p
    print(f"  creo el proyecto {TITULO_PROYECTO}")
    if SIMULAR:
        return None
    return gql("""mutation($o:ID!,$r:ID!,$t:String!){createProjectV2(input:{ownerId:$o,repositoryId:$r,title:$t}){
                  projectV2{id number title url}}}""",
               o=d["organization"]["id"], r=d["repository"]["id"], t=TITULO_PROYECTO)["createProjectV2"]["projectV2"]


def campos(pid):
    nodos = gql(f"query($p:ID!){{node(id:$p){{... on ProjectV2{{fields(first:50){{nodes{{{CAMPOS_PROYECTO}}}}}}}}}}}", p=pid)
    return {c["name"]: c for c in nodos["node"]["fields"]["nodes"] if c}


def opciones(lista):
    return "[" + ",".join(f"{{name:{q(n)},color:{c},description:{q(d)}}}" for n, c, d in lista) + "]"


def configurar_campos(pid, sprints):
    print("== Campos del proyecto")
    actuales = campos(pid)
    status = actuales["Status"]
    if [o["name"] for o in status["options"]] != [e[0] for e in ESTADOS]:
        print("  opciones de Status")
        gql(f"mutation{{updateProjectV2Field(input:{{fieldId:{q(status['id'])},singleSelectOptions:{opciones(ESTADOS)}}}){{projectV2Field{{... on ProjectV2FieldCommon{{id}}}}}}}}")
    nuevos = {"Tipo": ("SINGLE_SELECT", TIPOS), "Área": ("SINGLE_SELECT", AREAS), "Prioridad": ("SINGLE_SELECT", PRIORIDADES),
              "Wiegers": ("NUMBER", None), "Requisitos": ("TEXT", None), "Fecha": ("DATE", None), "Horas": ("NUMBER", None)}
    for nombre, (tipo, ops) in nuevos.items():
        if nombre in actuales:
            continue
        print(f"  creo el campo {nombre}")
        extra = f",singleSelectOptions:{opciones(ops)}" if ops else ""
        gql(f"mutation{{createProjectV2Field(input:{{projectId:{q(pid)},dataType:{tipo},name:{q(nombre)}{extra}}}){{clientMutationId}}}}")
    if "Sprint" not in actuales:
        print("  creo el campo Sprint (iteraciones)")
        iteraciones = ",".join(
            f"{{title:{q(f'Sprint {n}')},startDate:{q(str(s['inicio']))},duration:{(s['fin'] - s['inicio']).days + 1}}}"
            for n, s in sorted(sprints.items()))
        gql(f"""mutation{{createProjectV2Field(input:{{projectId:{q(pid)},dataType:ITERATION,name:"Sprint",
              iterationConfiguration:{{startDate:{q(str(sprints[1]['inicio']))},duration:7,iterations:[{iteraciones}]}}}}){{clientMutationId}}}}""")
    return campos(pid)


def items_proyecto(pid):
    items, cursor = {}, None
    while True:
        d = gql("""query($p:ID!,$c:String){node(id:$p){... on ProjectV2{items(first:100,after:$c){
                   pageInfo{hasNextPage endCursor}
                   nodes{id content{... on Issue{id}} fieldValueByName(name:"Status"){... on ProjectV2ItemFieldSingleSelectValue{name}}}}}}}""",
                p=pid, c=cursor)["node"]["items"]
        for it in d["nodes"]:
            if it["content"]:
                items[it["content"]["id"]] = {"id": it["id"], "status": (it["fieldValueByName"] or {}).get("name")}
        if not d["pageInfo"]["hasNextPage"]:
            return items
        cursor = d["pageInfo"]["endCursor"]


def estado_inicial(sprint, cerrado):
    if cerrado:
        return "Terminado"
    hoy = date.today()
    return "Por hacer" if sprint and SPRINTS[sprint]["inicio"] <= hoy else "Backlog"


def valor(campos_p, nombre, v):
    c = campos_p[nombre]
    if c["dataType"] == "SINGLE_SELECT":
        return f"{{singleSelectOptionId:{q(next(o['id'] for o in c['options'] if o['name'] == v))}}}"
    if c["dataType"] == "ITERATION":
        conf = c["configuration"]
        it = next(i for i in conf["iterations"] + conf["completedIterations"] if i["title"] == v)
        return f"{{iterationId:{q(it['id'])}}}"
    if c["dataType"] == "NUMBER":
        return f"{{number:{v}}}"
    if c["dataType"] == "DATE":
        return f"{{date:{q(str(v))}}}"
    return f"{{text:{q(v)}}}"


# ---------------------------------------------------------------- principal

def main():
    global SPRINTS
    SPRINTS, tareas = leer_excel()
    print(f"Planificación: {len(SPRINTS)} sprints, {len(tareas)} tareas{' (SIMULACIÓN)' if SIMULAR else ''}")
    horas_item = defaultdict(int)
    for t in tareas:
        if t["item"]:
            horas_item[t["item"]] += t["horas"]

    sincronizar_etiquetas()
    milestones = sincronizar_milestones(SPRINTS)
    existentes = issues_existentes()

    print("== Ítems del Product Backlog")
    nodo, campos_item = {}, {}
    for item in BACKLOG:
        codigo, tipo, titulo, reqs, prio, wiegers, sprint, area, asignados, _ = item
        etiquetas = {tipo, f"prioridad:{prio}", f"area:{area}"} | ({"opcional"} if sprint is None else set())
        cerrado = sprint is not None and SPRINTS[sprint]["fin"] < date.today()
        nodo[codigo] = asegurar_issue(existentes, codigo, f"{codigo} · {titulo}", cuerpo_item(item, horas_item[codigo]),
                                      etiquetas, asignados, milestones[sprint] if sprint else None, cerrado)
        estado_cerrado = cerrado or (existentes.get(codigo, {}).get("state") == "CLOSED")
        campos_item[codigo] = {"Tipo": TIPO_CAMPO[tipo], "Área": AREA_CAMPO[area],
                               "Prioridad": "Opcional" if sprint is None else prio.capitalize(),
                               "Requisitos": reqs, "Horas": horas_item[codigo],
                               **({"Wiegers": wiegers} if wiegers else {}),
                               **({"Sprint": f"Sprint {sprint}"} if sprint else {}),
                               "_estado": estado_inicial(sprint, estado_cerrado)}

    print("== Tareas")
    for t in tareas:
        cerrado = SPRINTS[t["sprint"]]["fin"] < date.today()
        nodo[t["codigo"]] = asegurar_issue(existentes, t["codigo"], t["titulo"], t["cuerpo"], {"tarea"}, [t["usuario"]],
                                           milestones[t["sprint"]], cerrado)
        estado_cerrado = cerrado or existentes.get(t["codigo"], {}).get("state") == "CLOSED"
        campos_item[t["codigo"]] = {"Tipo": "Tarea", "Sprint": f"Sprint {t['sprint']}", "Fecha": t["fecha"],
                                    "Horas": t["horas"], "_estado": estado_inicial(t["sprint"], estado_cerrado)}

    print("== Sub-issues (tarea -> ítem del backlog)")
    padre_actual = {c: (i["parent"] or {}).get("id") for c, i in existentes.items()}
    enlaces = [(nodo[t["item"]], nodo[t["codigo"]]) for t in tareas
               if t["item"] and nodo.get(t["codigo"]) and padre_actual.get(t["codigo"]) != nodo.get(t["item"])]
    print(f"  {len(enlaces)} enlaces por crear o corregir")
    if not SIMULAR:
        mutaciones_en_lote([f"addSubIssue(input:{{issueId:{q(p)},subIssueId:{q(h)},replaceParent:true}}){{clientMutationId}}"
                            for p, h in enlaces if p])

    print("== Tablero en GitHub Projects")
    proyecto = obtener_proyecto()
    if SIMULAR:
        print("  (simulación: no se configuran campos ni ítems)")
        return
    pid = proyecto["id"]
    gql("""mutation($p:ID!,$d:String!,$r:String!){updateProjectV2(input:{projectId:$p,shortDescription:$d,readme:$r}){clientMutationId}}""",
        p=pid, d="Product Backlog y Sprint Backlog de AulaConecta (Scrum + DevSecOps, 10 sprints)",
        r=("Ítems del **Product Backlog**: historias de usuario (HU), habilitadores técnicos (TEC) y entregables (DOC).\n\n"
           "Cada ítem tiene como sub-issues las tareas diarias de cada integrante (T{MMDD}-{M|A|C}), que forman el **Sprint Backlog**.\n\n"
           "Status: Backlog → Por hacer (planificación del sprint) → En progreso → En revisión → Terminado.\n\n"
           "Fuente: `docs/planificacion/Planificacion_v2.xlsx` · Script: `scripts/github/sincronizar_backlog.py`"))
    campos_p = configurar_campos(pid, SPRINTS)

    print("== Ítems del tablero")
    en_tablero = items_proyecto(pid)
    faltan = [c for c, n in nodo.items() if n and n not in en_tablero]
    print(f"  agrego {len(faltan)} issues")
    for i in range(0, len(faltan), 10):
        lote = faltan[i:i + 10]
        r = gql("mutation {\n" + "\n".join(
            f"m{j}: addProjectV2ItemById(input:{{projectId:{q(pid)},contentId:{q(nodo[c])}}}){{item{{id}}}}"
            for j, c in enumerate(lote)) + "\n}")
        for j, c in enumerate(lote):
            en_tablero[nodo[c]] = {"id": r[f"m{j}"]["item"]["id"], "status": None}
        time.sleep(1)

    print("== Valores de los campos")
    mutaciones = []
    for c, valores in campos_item.items():
        if not nodo.get(c):
            continue
        it = en_tablero[nodo[c]]
        valores = dict(valores)
        estado = valores.pop("_estado")
        if not it["status"]:
            valores["Status"] = estado
        for nombre, v in valores.items():
            mutaciones.append(f"updateProjectV2ItemFieldValue(input:{{projectId:{q(pid)},itemId:{q(it['id'])},"
                              f"fieldId:{q(campos_p[nombre]['id'])},value:{valor(campos_p, nombre, v)}}}){{clientMutationId}}")
    print(f"  {len(mutaciones)} valores")
    mutaciones_en_lote(mutaciones, tam=15)
    print(f"Listo: {proyecto['url']}")


if __name__ == "__main__":
    main()
